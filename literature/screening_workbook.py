"""03_Calisma/Eleme/Eleme_Calisma.xlsx: the title/abstract decisions file as a workbook for reading and checking.

The workbook is a view of eleme_kararlari.json. A file the program did not write last (or one that is open in Excel)
is never overwritten; a numbered copy is written next to it instead.
"""
import io
import uuid
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import screening
from .storage import digest, utcnow

FOLDER = screening.RELATIVE.parent.as_posix()
RELATIVE = f'{FOLDER}/Eleme_Calisma.xlsx'
BOLD = Font(bold=True)
HEAD = PatternFill('solid', fgColor='DDEBF7')
WRAP = Alignment(wrap_text=True, vertical='top')


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = BOLD, HEAD, WRAP
    for row in rows:
        ws.append(row)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def _number(record):
    try:
        return int(record.get('no'))
    except (TypeError, ValueError):
        return record.get('no', '')


def _order(record):
    number = _number(record)
    return (record.get('arama_konusu', ''), number if isinstance(number, int) else 10 ** 9, record.get('key', ''))


def build(data, criteria=''):
    """Workbook bytes; `criteria` is the text of eleme_olcutleri.md (optional)."""
    records, c = data['kayitlar'], screening.counts(data)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ozet'
    ws.append([f"Başlık-özet taraması ({data.get('proje', '')})".replace(' ()', '')])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    rows = [('Tekrarlar çıkarıldıktan sonra aday kayıt (01_Aday)', c['aday'])]
    if c['taranmadi']:
        rows.append(('Taranmayan (01_Aday içinde kararsız kalır)', c['taranmadi']))
    rows += [('Başlığı incelenen', c['baslik_incelenen']),
            ('Başlıktan hariç tutulan', c['baslik_haric']), ('Özeti okunan', c['ozet_incelenen']),
            ('Özetten hariç tutulan', c['ozet_haric'])] + ([('Özeti okunmayan', c['ozet_okunmadi'])] if c['ozet_okunmadi'] else []) + [
            ('Başlık-özet aşamasını geçen (02_BaslikOzet_Gecen)', c['gecen']),
            ('   Tam metin için seçilen', c['tam_metin_secilen']), ('   Yedek', c['yedek']),
            ('Başlık-özet aşamasında toplam hariç (05_Haric_BaslikOzet)', c['haric'])]
    rows += [(f'   {code}', n) for code, n in c['haric_kod'].items()]
    for row in rows:
        ws.append(list(row))
    total = c['haric'] + c['gecen'] + c['ozet_okunmadi'] + c['taranmadi']
    label = 'Kontrol: hariç + geçen + özeti okunmayan' + (' + taranmayan' if c['taranmadi'] else '') + ' = aday'
    ws.append([label, 'Doğru' if total == c['aday'] else 'HATALI'])
    ws.append(['Onay', screening.approval_state(data)])
    scope = data.get('kapsam') or {}
    if isinstance(scope, dict) and scope.get('tarama'):
        ws.append(['Tarama kapsamı', scope['tarama']])
        if scope.get('taranmayan'):
            ws.append(['', scope['taranmayan']])
    topics = sorted({r.get('arama_konusu', '') for r in records if r.get('arama_konusu')})
    if topics:
        ws.append([])
        ws.append(['Arama konusu', 'Geçen', 'Özetten hariç', 'Tam metin için seçilen'])
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = BOLD, HEAD
        for topic in topics:
            same = [r for r in records if r.get('arama_konusu') == topic]
            ws.append([topic, sum(r['durum'] == 'Geçti' for r in same), sum(r['durum'] == 'Hariç' for r in same),
                       sum(r.get('tam_metin') == 'Seçildi' for r in same)])
    selected = [r for r in records if r.get('tam_metin') == 'Seçildi']
    if selected:
        ws.append([])
        ws.append(['Tam metin için seçilenler'])
        ws[f'A{ws.max_row}'].font = BOLD
        kinds = Counter(r.get('tur', '') for r in selected)
        ws.append(['   Dergi makalesi', kinds['journalArticle']])
        ws.append(['   Bildiri', kinds['conferencePaper']])
        for year, n in sorted(Counter(str(r.get('yil', '')) for r in selected).items(), reverse=True):
            ws.append([f'   Yıl {year}', n])
        for question in ('AS1', 'AS2', 'AS3', 'AS4', 'AS5'):
            ws.append([f'   {question} ile ilgili', sum(question in str(r.get('AS', '')).split(';') for r in selected)])
        for access, n in sorted(Counter(r.get('erisim_on_kontrol', '') for r in selected).items()):
            ws.append([f'   Erişim: {access or "belirtilmemiş"}', n])
    ws.column_dimensions['A'].width = 62
    for column in 'BCD':
        ws.column_dimensions[column].width = 16

    _sheet(wb, f'TamMetin_{len(selected)}',
           ['No', 'Zotero anahtarı', 'Başlık', 'DOI', 'Tür', 'Yıl', 'Arama konusu', 'Araştırma sorusu', 'Seçim gerekçesi',
            'Erişim ön kontrolü', 'Özetten kanıt', 'Not'],
           [[_number(r), r['key'], r['baslik'], r.get('doi', ''), r.get('tur', ''), r.get('yil', ''), r.get('arama_konusu', ''),
             r.get('AS', ''), r.get('secim_gerekcesi', ''), r.get('erisim_on_kontrol', ''), r.get('kanit_alinti', ''),
             r.get('notlar', '')] for r in sorted(selected, key=_order)],
           [7, 12, 60, 26, 14, 6, 26, 16, 48, 20, 48, 40])
    spare = [r for r in records if r.get('tam_metin') == 'Yedek']
    _sheet(wb, 'Yedek',
           ['No', 'Zotero anahtarı', 'Başlık', 'DOI', 'Tür', 'Yıl', 'Arama konusu', 'Araştırma sorusu', 'Özetten kanıt', 'Not'],
           [[_number(r), r['key'], r['baslik'], r.get('doi', ''), r.get('tur', ''), r.get('yil', ''), r.get('arama_konusu', ''),
             r.get('AS', ''), r.get('kanit_alinti', ''), r.get('notlar', '')] for r in sorted(spare, key=_order)],
           [7, 12, 60, 26, 14, 6, 26, 16, 48, 40])
    excluded = [r for r in records if r['durum'] == 'Hariç']
    _sheet(wb, 'Haric_BaslikOzet',
           ['No', 'Zotero anahtarı', 'Başlık', 'DOI', 'Bakılan', 'Arama konusu', 'Gerekçe kodu', 'Gerekçe', 'Kanıt alıntısı', 'Not'],
           [[_number(r), r['key'], r['baslik'], r.get('doi', ''), r.get('bakilan', ''), r.get('arama_konusu', ''),
             r.get('gerekce_kodu', ''), r.get('gerekce', ''), r.get('kanit_alinti', ''), r.get('notlar', '')]
            for r in sorted(excluded, key=lambda r: (_order(r)[1], r['key']))],
           [7, 12, 60, 26, 9, 26, 8, 40, 48, 40])
    unread = [r for r in records if r['durum'] == 'Özet okunmadı']
    if unread:
        _sheet(wb, 'Ozet_Okunmadi', ['No', 'Zotero anahtarı', 'Başlık', 'DOI', 'Tür', 'Yıl'],
               [[_number(r), r['key'], r['baslik'], r.get('doi', ''), r.get('tur', ''), r.get('yil', '')]
                for r in sorted(unread, key=lambda r: (_order(r)[1], r['key']))], [7, 12, 60, 26, 14, 6])
    if criteria.strip():
        sheet = wb.create_sheet('Olcutler')
        sheet.column_dimensions['A'].width = 130
        for line in criteria.splitlines():
            heading = line.startswith('#')
            sheet.append([line.lstrip('# ') if heading else line])
            if heading:
                sheet[f'A{sheet.max_row}'].font = BOLD
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export(store):
    """Write the workbook; returns {'path', 'status'} like the project Excel."""
    data = screening.load(store.path(screening.RELATIVE.as_posix()))
    criteria_file = store.path(f'{FOLDER}/eleme_olcutleri.md')
    criteria = criteria_file.read_text(encoding='utf-8') if criteria_file.is_file() else ''
    content = build(data, criteria)
    relative = RELATIVE
    target = store.path(relative)
    with store.lock:
        last = store.db.execute("SELECT sha256 FROM export_runs WHERE kind='eleme_xlsx' AND relative_path=? ORDER BY rowid DESC LIMIT 1",
                                (relative,)).fetchone()
        ours = target.exists() and last and digest(target.read_bytes()) == last[0]
        open_in_excel = target.with_name('~$' + target.name).exists()
        status = 'updated'
        if target.exists() and (not ours or open_in_excel):
            relative = store._version_path(relative)
            status = 'versioned'
        try:
            store._atomic_write(relative, content, replace=status == 'updated')
        except PermissionError:
            relative = store._version_path(RELATIVE)
            store._atomic_write(relative, content)
            status = 'versioned'
        with store.db:
            store.db.execute('INSERT INTO export_runs VALUES(?,?,?,?,?,?)',
                             (str(uuid.uuid4()), utcnow(), 'eleme_xlsx', relative, digest(content), status))
    return {'path': str(store.path(relative)), 'status': status}
