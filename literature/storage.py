from __future__ import annotations

import hashlib
import html
import io
import json
import os
import re
import sqlite3
import threading
import textwrap
import uuid
import logging
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from zoneinfo import ZoneInfo

from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


class ValidationError(ValueError):
    pass


def utcnow():
    return datetime.now(timezone.utc).isoformat()


def digest(data: bytes):
    return hashlib.sha256(data).hexdigest()


def json_text(data):
    return json.dumps(data, ensure_ascii=False, separators=(',', ':'))


def safe_url(value: str):
    try:
        u = urlsplit(value)
        if u.scheme not in ('https', 'http') or not u.hostname:
            return ''
        # Search text has its own field. Do not archive session-bearing query strings.
        path = re.sub(r'(?i)/(token|session|auth)/[^/]+', r'/\1/[removed]', u.path)
        return urlunsplit((u.scheme, u.hostname, path, '', ''))
    except ValueError:
        return ''


# The first eleven columns keep their historical positions; new columns are appended.
SEARCH_HEADERS = ['Arama ID', 'Durum', 'Veritabanı / platform', 'Tarih', 'Tam sorgu',
                  'Alan', 'Filtre', 'Görünen sonuç sayısı', 'Dışa aktarılan kayıt sayısı',
                  'Ekran dosyası', 'Açıklama', 'Aktarım sayısının kaynağı', 'Export dosyaları',
                  'İşlemi yapan']
EXPORT_FILES_LIMIT = 2000
MATRIX_HEADERS = ['ID', 'IEEE künyesi', 'DOI', 'Yıl', 'Yayın türü', 'Kalite sınıfı ve yılı',
                  'Kalite kaynağı', 'Bulunma yolu / arama kaydı', 'Okuyan üye',
                  'Tamamlanan okuma geçişi', 'Okuma tarihi', 'Çekirdek adayı', 'Problem',
                  'Önerilen yöntem', 'Karşılaştırılan yöntemler', 'Veri kümesi / dil',
                  'Deney düzeni / donanım / bölünme', 'Metrikler', 'Sayısal ana sonuç ve PDF sayfası',
                  'Kaynak ihtiyacı / süre / maliyet', 'Açık kod / veri', 'Yazarın belirttiği sınırlılık',
                  'Tez açısından sınırlılık', 'Tema', 'Tezle ilişki', 'Yerel PDF dosyası',
                  'İndirilebilir PDF bağlantısı', 'Anotasyon dayanağı', 'İstisna gerekçesi']
# Full-text exclusions (Kılavuz s. 14, Hariç sayfası). Title/abstract exclusions are only counted.
EXCLUDED_HEADERS = ['Aday / Zotero anahtarı', 'Başlık', 'DOI veya bağlantı', 'Arama ID', 'Aşama',
                    'Gerekçe kodu', 'Açıklama', 'Karar veren', 'Karar tarihi']
ITEM_TYPES = {'journalArticle': 'Dergi makalesi', 'conferencePaper': 'Konferans bildirisi', 'preprint': 'Ön baskı',
              'bookSection': 'Kitap bölümü', 'book': 'Kitap', 'thesis': 'Tez', 'report': 'Rapor', 'webpage': 'Web sayfası'}

SCHEMA = '''
CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY, settings_json TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS capture_jobs(
 id TEXT PRIMARY KEY, context_json TEXT NOT NULL, state TEXT NOT NULL,
 created_at TEXT NOT NULL, captured_at TEXT, imported_at TEXT, image_path TEXT,
 image_json TEXT, extraction_json TEXT, prepared_json TEXT, error TEXT, run_id TEXT);
CREATE TABLE IF NOT EXISTS search_runs(
 id TEXT PRIMARY KEY, display_id TEXT NOT NULL UNIQUE, project_id TEXT NOT NULL,
 database_name TEXT NOT NULL, query_raw TEXT NOT NULL, query_verified INTEGER NOT NULL,
 fields TEXT, filters_json TEXT NOT NULL, count_raw TEXT, count_value INTEGER,
 count_qualifier TEXT, safe_url TEXT, observed_at_utc TEXT, timezone TEXT NOT NULL,
 operator TEXT, extraction_status TEXT, verification_status TEXT, notes TEXT,
 context_key TEXT, search_fingerprint TEXT, local_date TEXT, created_at TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS captures(
 id TEXT PRIMARY KEY, job_id TEXT NOT NULL UNIQUE, search_id TEXT NOT NULL REFERENCES search_runs(id),
 relative_path TEXT NOT NULL UNIQUE, sha256 TEXT NOT NULL, width INTEGER, height INTEGER,
 size_bytes INTEGER, source_mode TEXT, captured_at TEXT, imported_at TEXT, role TEXT,
 context_json TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS extraction_fields(
 job_id TEXT NOT NULL, field_name TEXT NOT NULL, value_json TEXT, source TEXT,
 adapter_version TEXT, observed_at TEXT, PRIMARY KEY(job_id,field_name));
CREATE TABLE IF NOT EXISTS audit_events(id INTEGER PRIMARY KEY, at TEXT, event TEXT, subject TEXT, details_json TEXT);
CREATE TABLE IF NOT EXISTS export_runs(id TEXT PRIMARY KEY, at TEXT, kind TEXT, relative_path TEXT, sha256 TEXT, status TEXT);
CREATE TABLE IF NOT EXISTS papers(
 id TEXT PRIMARY KEY, library_id TEXT NOT NULL, item_key TEXT NOT NULL, data_json TEXT NOT NULL,
 state TEXT NOT NULL DEFAULT 'candidate', annotation TEXT NOT NULL DEFAULT '',
 verified INTEGER NOT NULL DEFAULT 0, updated_at TEXT NOT NULL,
 UNIQUE(library_id,item_key));
CREATE TABLE IF NOT EXISTS search_paper_links(
 search_id TEXT REFERENCES search_runs(id), paper_id TEXT REFERENCES papers(id),
 PRIMARY KEY(search_id,paper_id));
CREATE TABLE IF NOT EXISTS reading_notes(
 paper_id TEXT PRIMARY KEY REFERENCES papers(id), zotero_json TEXT NOT NULL,
 fields_json TEXT NOT NULL, local_draft TEXT NOT NULL DEFAULT '');
CREATE TABLE IF NOT EXISTS screening_decisions(
 paper_id TEXT PRIMARY KEY REFERENCES papers(id), decision TEXT, reason TEXT, at TEXT);
CREATE TABLE IF NOT EXISTS search_refinements(
 id TEXT PRIMARY KEY, run_id TEXT UNIQUE REFERENCES search_runs(id), display_id TEXT NOT NULL,
 initial_count INTEGER NOT NULL, year_from INTEGER NOT NULL, year_to INTEGER NOT NULL,
 source_query TEXT NOT NULL, source_count INTEGER NOT NULL, query TEXT NOT NULL, reason TEXT NOT NULL,
 context_json TEXT NOT NULL, state TEXT NOT NULL, created_at TEXT NOT NULL, approved_at TEXT,
 child_run_id TEXT REFERENCES search_runs(id));
CREATE TABLE IF NOT EXISTS search_run_exports(
 run_id TEXT PRIMARY KEY REFERENCES search_runs(id), manual_count INTEGER, export_files TEXT NOT NULL DEFAULT '',
 zotero_count INTEGER, zotero_collection TEXT, zotero_checked_at TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS paper_search_codes(
 paper_id TEXT NOT NULL REFERENCES papers(id), code TEXT NOT NULL, PRIMARY KEY(paper_id,code));
CREATE TABLE IF NOT EXISTS paper_attachments(
 paper_id TEXT NOT NULL REFERENCES papers(id), attachment_key TEXT NOT NULL, filename TEXT NOT NULL DEFAULT '',
 content_type TEXT NOT NULL DEFAULT '', link_mode TEXT NOT NULL DEFAULT '', PRIMARY KEY(paper_id,attachment_key));
'''


class Store:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.lock = threading.RLock()
        for folder in ('02_Ekranlar', '03_Calisma/Veri/gecici', '05_Teslim', 'Yedekler'):
            self.path(folder).mkdir(parents=True, exist_ok=True)
        self.db_path = self.path('03_Calisma/Veri/kayit.sqlite3')
        self.db = sqlite3.connect(self.db_path, check_same_thread=False, timeout=15)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('PRAGMA synchronous=FULL')
        version = self.db.execute('PRAGMA user_version').fetchone()[0]
        if version > 2:
            self.db.close()
            raise ValidationError('Bu proje daha yeni bir uygulama sürümü gerektiriyor.')
        if version == 1:
            # SQLite backup includes the WAL; copying the main file alone would lose recent captures.
            backup = self.path('Yedekler/schema1_' + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S') + '_' + os.urandom(4).hex() + '.sqlite3')
            target = sqlite3.connect(backup)
            try:self.db.backup(target)
            finally:target.close()
        from .workflow import WORKFLOW_SCHEMA
        from .zotero_writer import WRITER_SCHEMA
        try:self.db.executescript('BEGIN IMMEDIATE;\n' + SCHEMA + WORKFLOW_SCHEMA + WRITER_SCHEMA + '\nPRAGMA user_version=2;\nCOMMIT;')
        except Exception:
            self.db.rollback();self.db.close();raise
        with self.db:
            self.db.execute("UPDATE workflow_sessions SET state='paused',error='Uygulama yeniden açıldı. Kaldığınız adımı kontrol edip devam edin.' WHERE state='running'")
        if not self.db.execute('SELECT 1 FROM projects').fetchone():
            with self.db:
                self.db.execute('INSERT INTO projects VALUES (?,?)', (str(uuid.uuid4()), json_text({
                    'team_code': 'OEC-1017', 'project_name': 'Literatür taraması',
                    'operator': os.environ.get('USERNAME', ''), 'timezone': 'Europe/Istanbul',
                    'hotkey': 'CTRL+M', 'monitor': 'active', 'clock_confirmed': False,
                    'collection_mapping': {}, 'library': 'users/0', 'zotero_mode': 'local'})))
        self.recover()

    def close(self):
        self.db.close()

    def path(self, relative):
        if not isinstance(relative, str) or not relative or ':' in relative:
            raise ValidationError('Geçersiz proje yolu.')
        raw = Path(relative)
        if raw.is_absolute() or '..' in raw.parts:
            raise ValidationError('Dosya proje klasörü dışında olamaz.')
        result = (self.root / raw).resolve()
        if not result.is_relative_to(self.root):
            raise ValidationError('Bağlantı proje klasörü dışına çıkıyor.')
        return result

    def settings(self):
        with self.lock:
            return json.loads(self.db.execute('SELECT settings_json FROM projects').fetchone()[0])

    def update_settings(self, updates):
        allowed = {'team_code', 'project_name', 'operator', 'hotkey', 'monitor',
                   'clock_confirmed', 'collection_mapping', 'library', 'zotero_mode', 'zotero_data_dir'}
        if not isinstance(updates, dict) or set(updates) - allowed:
            raise ValidationError('Bilinmeyen ayar.')
        if 'team_code' in updates and not re.fullmatch(r'[A-Za-z0-9_-]{1,40}', updates['team_code']):
            raise ValidationError('Takım kodu yalnız harf, sayı, tire ve alt çizgi içerebilir.')
        if 'library' in updates and not re.fullmatch(r'(users|groups)/\d+', updates['library']):
            raise ValidationError('Zotero kütüphanesi users/ID veya groups/ID olmalı.')
        if 'zotero_data_dir' in updates and (not isinstance(updates['zotero_data_dir'], str) or len(updates['zotero_data_dir']) > 500):
            raise ValidationError('Zotero veri klasörü yolu geçersiz.')
        if len(json_text(updates)) > 16000:
            raise ValidationError('Ayarlar çok uzun.')
        with self.lock, self.db:
            settings = self.settings()
            if 'team_code' in updates and updates['team_code'] != settings['team_code']:
                settings.pop('managed_excel',None)
            settings.update(updates)
            self.db.execute('UPDATE projects SET settings_json=?', (json_text(settings),))
        return settings

    def audit(self, event, subject, details=None):
        self.db.execute('INSERT INTO audit_events(at,event,subject,details_json) VALUES(?,?,?,?)',
                        (utcnow(), event, subject, json_text(details or {})))

    def begin_job(self, job_id, context):
        if not re.fullmatch(r'[\w-]{1,80}', job_id):
            raise ValidationError('Geçersiz iş kimliği.')
        with self.lock, self.db:
            self.db.execute('INSERT OR IGNORE INTO capture_jobs(id,context_json,state,created_at) VALUES(?,?,?,?)',
                            (job_id, json_text(context), 'capturing', utcnow()))
        return self.get_job(job_id)

    def get_job(self, job_id):
        with self.lock:
            row = self.db.execute('SELECT * FROM capture_jobs WHERE id=?', (job_id,)).fetchone()
        if not row:
            raise ValidationError('Yakalama işi bulunamadı.')
        return dict(row)

    def save_capture(self, job_id, png, captured_at=None, source='windows'):
        with self.lock:
            old = self.get_job(job_id)
            if old['image_path']:
                return old
            if len(png) > 96 * 1024 * 1024:
                raise ValidationError('PNG boyutu sınırı aşıldı.')
            with Image.open(io.BytesIO(png)) as im:
                if im.format != 'PNG':
                    raise ValidationError('Yalnız özgün PNG kabul edilir.')
                width, height = im.size
                im.verify()
            relative = f'03_Calisma/Veri/gecici/{job_id}.png'
            info = dict(sha256=digest(png), width=width, height=height, size_bytes=len(png), source_mode=source)
            captured_at = captured_at or (utcnow() if source == 'windows' else None)
            imported_at = utcnow() if source == 'import' else None
            # Journal the intended file before writing it, so startup can reconcile a crash.
            with self.db:
                self.db.execute('UPDATE capture_jobs SET image_path=?,image_json=?,captured_at=?,imported_at=? WHERE id=?',
                                (relative, json_text(info), captured_at, imported_at, job_id))
            try:
                self._atomic_write(relative, png)
            except OSError:
                self.mark_incomplete(job_id, 'Görüntü diske yazılamadı; boş alanı ve klasör iznini kontrol edin.')
                raise
            with self.db:
                self.db.execute("UPDATE capture_jobs SET state='file_saved',error=NULL WHERE id=?", (job_id,))
            return self.get_job(job_id)

    def _atomic_write(self, relative, data, replace=False):
        target = self.path(relative)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(target.name + '.' + uuid.uuid4().hex + '.tmp')
        try:
            with temporary.open('xb') as f:
                f.write(data)
                f.flush()
                os.fsync(f.fileno())
            if replace:
                os.replace(temporary, target)
            else:
                # On Windows rename fails if the destination exists. Never replace original evidence.
                if target.exists():
                    raise FileExistsError(target)
                os.rename(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)

    def mark_incomplete(self, job_id, reason, extraction=None):
        with self.lock, self.db:
            self.db.execute("UPDATE capture_jobs SET state='incomplete',error=?,extraction_json=COALESCE(?,extraction_json) WHERE id=? AND state!='verified'",
                            (reason, json_text(extraction) if extraction is not None else None, job_id))
        return self.get_job(job_id)

    def complete_job(self, job_id, data, context, *, manual=False, new_run=False, attach_to=None, role='sonuc'):
        with self.lock:
            job = self.get_job(job_id)
            committed=self.db.execute('SELECT search_id FROM captures WHERE job_id=?',(job_id,)).fetchone()
            if committed:
                return self.get_run(committed[0])
            if job['state'] == 'verified':
                return self.get_run(job['run_id'])
            if not job['image_path'] or not self.path(job['image_path']).exists():
                return self.mark_incomplete(job_id, 'Görüntü dosyası bulunamadı.', data)
            if digest(self.path(job['image_path']).read_bytes()) != json.loads(job['image_json'])['sha256']:
                return self.mark_incomplete(job_id, 'PNG bütünlük kontrolü başarısız; dosya sonradan değişmiş.', data)
            if not manual and json.loads(job['context_json']) != context:
                return self.mark_incomplete(job_id, 'Yakalama sırasında sekme veya sayfa değişti. Görüntüyü kontrol edin.', data)
            required = bool(data.get('database')) and bool(str(data.get('query_raw', '')).strip())
            count = data.get('count_value')
            valid_count = type(count) is int and 0 <= count <= 2**53 - 1
            if not required or not valid_count or not data.get('query_verified') or data.get('status') != 'complete':
                return self.mark_incomplete(job_id, 'Sorgu/sonuç doğrulanamadı. Eksik alanları tamamlayın.', data)
            if len(str(data['query_raw'])) > 32767:
                return self.mark_incomplete(job_id, 'Sorgu Excel hücre sınırını aşıyor; otomatik aktarılmadı.', data)
            if role not in ('sonuc', 'sorgu', 'filtre'):
                raise ValidationError('Geçersiz kanıt türü.')
            data = dict(data, safe_url=safe_url(data.get('safe_url', '')))
            now = utcnow()
            tz = ZoneInfo('Europe/Istanbul')
            observed = data.get('observed_at_utc') or now
            # Unknown import capture times stay unknown; use a clearly marked import date in filenames.
            stamp = job['captured_at'] or job['imported_at'] or now
            local_date = datetime.fromisoformat(stamp).astimezone(tz).strftime('%Y%m%d')
            signature = digest(json_text([data['database'], data['query_raw'], data.get('fields', ''),
                                          data.get('filters', []), count, data.get('count_qualifier')]).encode())
            context_key = digest(json_text(context).encode())
            row = None
            if attach_to:
                row = self.get_run(attach_to)
                if row['search_fingerprint'] != signature:
                    raise ValidationError('Sorgu veya filtreler farklı; yeni arama olarak kaydedin.')
            elif not new_run and not manual:
                row = self.db.execute('SELECT * FROM search_runs WHERE context_key=? AND search_fingerprint=? AND local_date=? ORDER BY created_at DESC LIMIT 1',
                                      (context_key, signature, local_date)).fetchone()
            if row:
                run = dict(row)
            else:
                numbers = [int(r[0][1:]) for r in self.db.execute('SELECT display_id FROM search_runs') if re.fullmatch(r'S\d+', r[0])]
                display = f'S{max(numbers, default=0)+1:03d}'
                settings = self.settings()
                run = dict(id=str(uuid.uuid4()), display_id=display,
                           project_id=self.db.execute('SELECT id FROM projects').fetchone()[0],
                           database_name=data['database'], query_raw=data['query_raw'], query_verified=1,
                           fields=data.get('fields', ''), filters_json=json_text(data.get('filters', [])),
                           count_raw=data.get('count_raw', str(count)), count_value=count,
                           count_qualifier=data.get('count_qualifier', 'exact'), safe_url=data['safe_url'],
                           observed_at_utc=observed, timezone='Europe/Istanbul', operator=settings['operator'],
                           extraction_status='manual' if manual else 'automatic',
                           verification_status='user_confirmed' if manual else 'context_checked',
                           notes=data.get('notes', ''), context_key=context_key, search_fingerprint=signature,
                           local_date=local_date, created_at=now, updated_at=now)
            prefix = {'sonuc': 'a', 'sorgu': 'b', 'filtre': 'c'}[role]
            import_tag = '_iceaktarim' if not job['captured_at'] else ''
            stem = f"{run['display_id']}_{local_date}{import_tag}_{prefix}_{role}"
            directory = f"02_Ekranlar/{run['display_id']}"
            version = 1
            relative = f'{directory}/{stem}.png'
            while self.path(relative).exists() or self.db.execute('SELECT 1 FROM captures WHERE relative_path=?', (relative,)).fetchone():
                version += 1
                relative = f'{directory}/{stem}_v{version:02d}.png'
            image = json.loads(job['image_json'])
            prepared = dict(run=run, capture=dict(id=str(uuid.uuid4()), job_id=job_id, search_id=run['id'],
                relative_path=relative, **image, captured_at=job['captured_at'], imported_at=job['imported_at'],
                role=role, context_json=job['context_json']))
            with self.db:
                self.db.execute("UPDATE capture_jobs SET state='committing',prepared_json=?,extraction_json=? WHERE id=?",
                                (json_text(prepared), json_text(data), job_id))
            self.path(directory).mkdir(parents=True, exist_ok=True)
            os.rename(self.path(job['image_path']), self.path(relative))
            self._finish_commit(job_id, prepared, data)
            try:self._write_gallery(run['id'])
            except OSError:
                logging.warning('Evidence saved; gallery could not be written for %s',run['display_id'])
            return self.get_run(run['id'])

    def _finish_commit(self, job_id, prepared, extraction):
        run, capture = prepared['run'], prepared['capture']
        with self.db:
            self.db.execute(f"INSERT OR IGNORE INTO search_runs({','.join(run)}) VALUES({','.join('?' for _ in run)})", tuple(run.values()))
            self.db.execute(f"INSERT OR IGNORE INTO captures({','.join(capture)}) VALUES({','.join('?' for _ in capture)})", tuple(capture.values()))
            for key in ('query_raw', 'fields', 'filters', 'count_raw', 'count_value', 'count_qualifier'):
                self.db.execute('INSERT OR REPLACE INTO extraction_fields VALUES(?,?,?,?,?,?)',
                                (job_id, key, json_text(extraction.get(key)), extraction.get('source', 'manual'),
                                 extraction.get('adapter_version', ''), extraction.get('observed_at_utc', utcnow())))
            self.db.execute("UPDATE capture_jobs SET state='verified',run_id=?,image_path=?,error=NULL WHERE id=?",
                            (run['id'], capture['relative_path'], job_id))
            self.audit('capture_saved', job_id, {'search': run['display_id']})

    def recover(self):
        with self.lock:
            jobs = list(self.db.execute("SELECT * FROM capture_jobs WHERE state!='verified'"))
            for raw in jobs:
                job = dict(raw)
                if job['prepared_json']:
                    prepared = json.loads(job['prepared_json'])
                    final = self.path(prepared['capture']['relative_path'])
                    if final.exists() and digest(final.read_bytes()) == prepared['capture']['sha256']:
                        self._finish_commit(job['id'], prepared, json.loads(job['extraction_json'] or '{}'))
                        try:self._write_gallery(prepared['run']['id'])
                        except OSError:logging.warning('Recovered evidence; gallery unavailable')
                        continue
                self.mark_incomplete(job['id'], 'Önceki işlem tamamlanamadı. Kaydedilmiş görüntüyü kontrol edin.')

    def get_run(self, identity):
        row = self.db.execute('SELECT * FROM search_runs WHERE id=? OR display_id=?', (identity, identity)).fetchone()
        if not row:
            raise ValidationError('Arama bulunamadı.')
        return dict(row)

    def list_runs(self):
        with self.lock:
            return [dict(r) for r in self.db.execute('SELECT * FROM search_runs ORDER BY created_at DESC')]

    def captures(self, search_id):
        with self.lock:
            return [dict(r) for r in self.db.execute('SELECT * FROM captures WHERE search_id=? ORDER BY rowid', (search_id,))]

    def pending(self):
        with self.lock:
            return [dict(r) for r in self.db.execute("SELECT * FROM capture_jobs WHERE state!='verified' ORDER BY created_at DESC")]

    def update_run(self, identity, changes):
        """Edit the user's description and export bookkeeping; evidence fields stay immutable."""
        if not isinstance(changes, dict) or not changes or set(changes) - {'notes', 'manual_count', 'export_files'}:
            raise ValidationError('Yalnız açıklama ve dışa aktarım bilgisi güncellenebilir; kanıt alanları yeni doğrulama gerektirir.')
        if 'notes' in changes and (not isinstance(changes['notes'], str) or len(changes['notes']) > 30000):
            raise ValidationError('Açıklama çok uzun.')
        manual = changes.get('manual_count')
        if 'manual_count' in changes:
            if manual in ('', None):
                manual = None
            elif isinstance(manual, str) and manual.strip().isdigit():
                manual = int(manual.strip())
            if manual is not None and (type(manual) is not int or not 0 <= manual <= 2**53 - 1):
                raise ValidationError('Aktarılan kayıt sayısı 0 veya daha büyük bir tam sayı olmalı.')
        files = changes.get('export_files')
        if 'export_files' in changes and (not isinstance(files, str) or len(files) > EXPORT_FILES_LIMIT):
            raise ValidationError(f'Export dosya adları en fazla {EXPORT_FILES_LIMIT} karakter olabilir.')
        with self.lock, self.db:
            run = self.get_run(identity)
            now = utcnow()
            if 'notes' in changes:
                self.db.execute('UPDATE search_runs SET notes=?,updated_at=? WHERE id=?', (changes['notes'], now, run['id']))
                self.audit('note_edited', run['id'])
            if 'manual_count' in changes or 'export_files' in changes:
                self.db.execute('INSERT OR IGNORE INTO search_run_exports(run_id,updated_at) VALUES(?,?)', (run['id'], now))
                if 'manual_count' in changes:
                    self.db.execute('UPDATE search_run_exports SET manual_count=?,updated_at=? WHERE run_id=?', (manual, now, run['id']))
                if 'export_files' in changes:
                    self.db.execute('UPDATE search_run_exports SET export_files=?,updated_at=? WHERE run_id=?', (files.strip(), now, run['id']))
                self.audit('export_info_edited', run['id'], {'manual_count': manual if 'manual_count' in changes else 'unchanged'})

    def export_info(self, run_id):
        """Return export bookkeeping for one search: manual entry and the latest Zotero collection count."""
        row = self.db.execute('SELECT * FROM search_run_exports WHERE run_id=?', (run_id,)).fetchone()
        info = dict(row) if row else {'run_id': run_id, 'manual_count': None, 'export_files': '',
                                      'zotero_count': None, 'zotero_collection': None, 'zotero_checked_at': None}
        workflow = self.db.execute('SELECT id FROM workflow_steps WHERE run_id=?', (run_id,)).fetchone()
        info['ris_count'] = None
        if workflow:
            info['ris_count'] = self.db.execute("SELECT COALESCE(SUM(record_count),0) FROM workflow_exports WHERE step_id=? AND state='complete'",
                                                (workflow['id'],)).fetchone()[0]
        if info['ris_count'] is not None:
            info['effective_count'], info['source'] = info['ris_count'], 'Otomatik RIS doğrulaması'
        elif info['manual_count'] is not None:
            info['effective_count'], info['source'] = info['manual_count'], 'Elle girildi'
        elif info['zotero_count'] is not None:
            info['effective_count'], info['source'] = info['zotero_count'], f"Zotero koleksiyonu {info['zotero_collection'] or ''}".strip()
        else:
            info['effective_count'], info['source'] = None, ''
        if (info['manual_count'] is not None and info['zotero_count'] is not None
                and info['manual_count'] != info['zotero_count']):
            info['source'] += f" (Zotero'da {info['zotero_count']} kayıt; farkı kontrol edin)"
        return info

    def _write_gallery(self, search_id):
        run = self.get_run(search_id)
        items = []
        for cap in self.captures(search_id):
            name = html.escape(Path(cap['relative_path']).name, quote=True)
            when = cap['captured_at'] or 'Çekim zamanı bilinmiyor; elle içe aktarıldı'
            items.append(f'<article><h2>{name}</h2><p>Uygulamanın kayıt bilgisi: {html.escape(when)}</p>'
                         f'<p>SHA-256: {cap["sha256"]}</p><a href="{name}">Özgün PNG</a><br><img src="{name}" alt="Özgün ekran görüntüsü"></article>')
        page = ('<!doctype html><html lang="tr"><meta charset="utf-8"><meta name="viewport" content="width=device-width">'
                '<title>Arama kanıtları</title><style>body{font:16px system-ui;max-width:1300px;margin:24px auto;padding:20px}img{max-width:100%}article{margin-bottom:48px}p{overflow-wrap:anywhere}</style>'
                f'<h1>{run["display_id"]} — Arama kanıtları</h1><p>{html.escape(run["query_raw"])}</p>' + ''.join(items) + '</html>')
        self._atomic_write(f"02_Ekranlar/{run['display_id']}/index.html", page.encode('utf-8'), replace=True)

    @staticmethod
    def _cell(ws, row, col, value):
        cell = ws.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = 's'
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell

    @staticmethod
    def _format_excel_sheet(ws):
        widths = {
            'Arama_Kaydi': [10, 22, 18, 22, 58, 18, 36, 16, 18, 33, 36, 26, 34, 22],
            'Matris': [12, 52, 30, 12, 20, 23, 30, 26, 20, 24, 18, 18,
                       44, 44, 38, 32, 40, 30, 44, 36, 26, 44, 44, 26, 38, 32, 36, 36, 36],
            'Hariç': [16, 52, 36, 16, 16, 14, 52, 24, 16],
        }[ws.title]
        ws.freeze_panes = 'B2'
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:1'
        ws.auto_filter.ref = ws.dimensions
        ws.print_options.horizontalCentered = True
        header = PatternFill('solid', fgColor='163D4A')
        stripe = PatternFill('solid', fgColor='F0F5F7')
        line = Border(bottom=Side(style='hair', color='DCE5E9'))
        for cell, width in zip(ws[1], widths):
            ws.column_dimensions[cell.column_letter].width = width
            cell.fill = header
            cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 46
        for row in ws.iter_rows(min_row=2):
            lines = 1
            for cell, width in zip(row, widths):
                linked = bool(cell.hyperlink)
                cell.font = Font(name='Calibri', size=11, color='0563C1' if linked else '243B47',
                                 bold=cell.column == 1, underline='single' if linked else None)
                cell.border = line
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = stripe
                # Excel does not auto-fit wrapped rows when opening generated files.
                paragraphs = str(cell.value or '').split('\n')
                lines = max(lines, sum(max(1, len(textwrap.wrap(p, max(1, int(width) - 3)))) for p in paragraphs))
            ws.row_dimensions[row[0].row].height = min(409, max(32, lines * 15 + 12))

    def export_excel(self):
        with self.lock:
            settings = self.settings()
            name = f"{settings['team_code']}_O2_LiteraturMatrisi"
            relative = settings.get('managed_excel', f'03_Calisma/{name}.xlsx')
            target = self.path(relative)
            last = self.db.execute("SELECT sha256 FROM export_runs WHERE kind='xlsx' AND relative_path=? ORDER BY rowid DESC LIMIT 1", (relative,)).fetchone()
            conflict = target.exists() and (not last or digest(target.read_bytes()) != last[0])
            if conflict:
                # Never overwrite an existing user workbook, including the initial empty template.
                backup = f'Yedekler/{name}_onceki_{datetime.now().strftime("%Y%m%d_%H%M%S")}_{uuid.uuid4().hex[:6]}.xlsx'
                self._atomic_write(backup, target.read_bytes())
            wb = Workbook()
            wb.properties.title = settings['project_name']
            wb.properties.creator = settings['operator']
            wb.properties.subject = settings['team_code']
            wb.properties.description = 'SQLite kayıtlarının görünümü. Düzenleme uygulamadan yapılır.'
            wb.remove(wb.active)
            for title, headers in [('Arama_Kaydi', SEARCH_HEADERS), ('Matris', MATRIX_HEADERS), ('Hariç', EXCLUDED_HEADERS)]:
                ws = wb.create_sheet(title)
                ws.append(headers)
            ws = wb['Arama_Kaydi']
            for row_index, run in enumerate(reversed(self.list_runs()), 2):
                self._write_gallery(run['id'])
                caps = self.captures(run['id'])
                count = run['count_value'] if run['count_qualifier'] == 'exact' else run['count_raw']
                info = self.export_info(run['id'])
                when = datetime.fromisoformat(run['observed_at_utc']).astimezone(ZoneInfo('Europe/Istanbul')).replace(tzinfo=None)
                values = [run['display_id'], 'Kullanıcı doğruladı' if run['verification_status']=='user_confirmed' else 'Bağlam kontrol edildi',
                          run['database_name'], when, run['query_raw'], run['fields'],
                          '\n'.join(json.loads(run['filters_json'])), count, info['effective_count'],
                          '\n'.join(Path(c['relative_path']).name for c in caps), run['notes'],
                          info['source'], info['export_files'], run['operator'] or '']
                for column, value in enumerate(values, 1):
                    self._cell(ws, row_index, column, value)
                ws.cell(row_index, 4).number_format = 'dd.mm.yyyy hh:mm:ss'
                ws.cell(row_index, 8).number_format = '#,##0'
                ws.cell(row_index, 9).number_format = '#,##0'
                ws.cell(row_index, 10).hyperlink = f"../02_Ekranlar/{run['display_id']}/index.html"
            self._export_papers(wb)
            for sheet in wb:
                self._format_excel_sheet(sheet)
            buffer = io.BytesIO()
            wb.save(buffer)
            wb.close()
            data = buffer.getvalue()
            if conflict:
                relative = self._version_path(f'03_Calisma/{name}.xlsx')
            try:
                # Some Excel installations permit replacement despite an open workbook.
                if not conflict and target.with_name('~$' + target.name).exists():
                    raise PermissionError('Excel çalışma kitabı açık.')
                self._atomic_write(relative, data, replace=not conflict)
                status = 'versioned' if conflict else 'updated'
            except PermissionError:
                relative = self._version_path(f'03_Calisma/{name}.xlsx')
                self._atomic_write(relative, data)
                status = 'locked_versioned'
            with self.db:
                self.db.execute('INSERT INTO export_runs VALUES(?,?,?,?,?,?)',
                                (str(uuid.uuid4()), utcnow(), 'xlsx', relative, digest(data), status))
                settings['managed_excel'] = relative
                self.db.execute('UPDATE projects SET settings_json=?', (json_text(settings),))
            return {'path': str(self.path(relative)), 'relative_path': relative, 'status': status}

    def _version_path(self, relative):
        p = Path(relative)
        index = 2
        while self.path(str(p.with_name(f'{p.stem}_v{index:02d}{p.suffix}'))).exists():
            index += 1
        return str(p.with_name(f'{p.stem}_v{index:02d}{p.suffix}')).replace('\\', '/')

    def paper_search_codes(self, paper_id):
        """S codes linked by hand in the manager plus those taken from Zotero S collections."""
        manual = [r[0] for r in self.db.execute('SELECT s.display_id FROM search_runs s JOIN search_paper_links l ON l.search_id=s.id WHERE l.paper_id=?', (paper_id,))]
        zotero = [r[0] for r in self.db.execute('SELECT code FROM paper_search_codes WHERE paper_id=?', (paper_id,))]
        return sorted(set(manual) | set(zotero), key=lambda c: (len(c), c))

    def paper_rows(self, state):
        """Papers of one state with parsed note fields, in matrix order (L001, L002 … then unnumbered)."""
        rows = []
        for paper in self.db.execute('SELECT * FROM papers WHERE state=?', (state,)):
            note = self.db.execute('SELECT fields_json FROM reading_notes WHERE paper_id=?', (paper['id'],)).fetchone()
            rows.append((dict(paper), json.loads(paper['data_json']), json.loads(note[0]) if note else {}))
        def order(row):
            match = re.fullmatch(r'L(\d+)', row[2].get('Matris ID', '').strip())
            return (0, int(match.group(1)), '') if match else (1, 0, row[0]['item_key'])
        return sorted(rows, key=order)

    def _export_papers(self, wb):
        from .pdf_archive import archive_name
        ws = wb['Matris']
        for paper, data, fields in self.paper_rows('included'):
            year = re.search(r'\b(19|20)\d{2}\b', str(data.get('date', '')))
            kind = fields.get('Yayın türü', '').strip() or ITEM_TYPES.get(data.get('itemType', ''), data.get('itemType', ''))
            values = [fields.get('Matris ID', '').strip() or paper['item_key'], data.get('citation') or data.get('title', ''),
                      data.get('DOI', ''), year.group(0) if year else data.get('date', ''), kind,
                      fields.get('Kalite sınıfı ve yılı', ''), fields.get('Kalite kaynağı', ''),
                      ', '.join(self.paper_search_codes(paper['id'])),
                      fields.get('Okuyan kişi', ''), fields.get('Okuma düzeyi', ''), fields.get('Okuma tarihi', ''), fields.get('Çekirdek adayı', '')]
            values += [fields.get(k, '') for k in ['Amaç/problem', 'Yöntem', 'Karşılaştırılan yöntemler', 'Veri seti ve dil',
                       'Deney düzeni', 'Değerlendirme metrikleri', 'Temel bulgu', 'Kaynak ihtiyacı', 'Açık kod/veri',
                       'Yazarların sınırlılıkları', 'Benim değerlendirmem', 'Tema', 'Tezle ilişki']]
            values += [archive_name(fields, data), fields.get('PDF bağlantısı', ''), fields.get('Kanıt sayfaları', ''),
                       fields.get('İstisna gerekçesi', '')]
            row = ws.max_row + 1
            for col, value in enumerate(values, 1):
                self._cell(ws, row, col, value)
        ws = wb['Hariç']
        for paper, data, fields in self.paper_rows('excluded'):
            decision = self.db.execute('SELECT at FROM screening_decisions WHERE paper_id=?', (paper['id'],)).fetchone()
            explanation = fields.get('Dahil/hariç gerekçesi', '')
            code = fields.get('Gerekçe kodu', '').strip()
            if not code:
                match = re.search(r'\bH[1-6]\b', explanation)
                code = match.group(0) if match else ''
            when = fields.get('Karar tarihi', '').strip()
            if not when and decision and decision['at']:
                when = datetime.fromisoformat(decision['at']).astimezone(ZoneInfo('Europe/Istanbul')).strftime('%d.%m.%Y')
            values = [paper['item_key'], data.get('title', ''), data.get('DOI') or data.get('url', ''),
                      ', '.join(self.paper_search_codes(paper['id'])), fields.get('Aşama', '').strip() or 'Tam metin',
                      code, explanation, fields.get('Karar veren', ''), when]
            row = ws.max_row + 1
            for col, value in enumerate(values, 1):
                self._cell(ws, row, col, value)
