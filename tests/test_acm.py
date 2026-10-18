import csv
import io
import os
import time
import unittest
from unittest.mock import patch

import test_storage as fixtures
from test_zotero_writer import FakeLocalAPI
from literature import acm, ham_export
from literature.storage import ValidationError
from literature.zotero_writer import ZoteroWriter

QUERY = ('[[Abstract: "retrieval augmented generation"] OR [Abstract: "retrieval-augmented generation"]] AND '
         '[[Abstract: turkish] OR [Abstract: multilingual]] AND [E-Publication Date: (01/01/2024 TO 12/31/2026)]')
CONTEXT = dict(profile='p', windowId=1, tabId=2, documentId='d', epoch=0)


def item(i, kind='research-article', doi=None):
    return {'doi': doi or f'10.1145/3805712.{3800000 + i}', 'type': kind, 'published': '19 July 2026',
            'title': f'Paper {i}', 'authors': ['Rui Qi', 'Jian-Yun Nie'],
            'venue': "SIGIR '26: Proceedings of the 49th International ACM SIGIR Conference", 'pages': '4045–4051',
            'page': 1 + i // 50, 'position': 1 + i % 50}


CROSSREF = {'type': 'proceedings-article', 'title': ['Crossref Title'], 'container-title': ['Proceedings of SIGIR'],
            'author': [{'given': 'Rui', 'family': 'Qi'}, {'given': 'Jian-Yun', 'family': 'Nie'}],
            'issued': {'date-parts': [[2026, 7, 13]]}, 'page': '4045-4051', 'publisher': 'ACM',
            'ISBN': ['9798400723537'], 'event': {'name': 'SIGIR 2026', 'location': 'Melbourne, Australia'},
            'abstract': '<jats:title>Abstract</jats:title><jats:p>A multilingual collection.</jats:p>'}


class FakeWriter:
    def __init__(self, error=None):
        self.calls, self.error = [], error

    def import_acm(self, records, search_id):
        self.calls.append((records, search_id))
        if self.error:
            raise self.error
        return {'imported': len(records), 'reused': 0, 'failed': 0, 'failures': []}


class AcmTest(unittest.TestCase):
    setUp = fixtures.StorageTest.setUp
    tearDown = fixtures.StorageTest.tearDown

    def run_(self, job='a', count=4, database=acm.DATABASE, filters=None):
        self.store.begin_job(job, CONTEXT)
        self.store.save_capture(job, self.png.read_bytes())
        data = fixtures.fields(database=database, query_raw=QUERY, fields='Abstract', count_value=count,
                               count_raw=f'{count} Results',
                               filters=filters if filters is not None else ['Yıl: 2024–2026', acm.TYPE_FILTER])
        return self.store.complete_job(job, data, CONTEXT)

    def collection(self, items):
        return {'items': items, 'pages': 1, 'total': len(items), 'query': QUERY, 'url': 'https://dl.acm.org/action/doSearch'}

    def offered(self, count=4):
        run = self.run_(count=count)
        return run, acm.propose(self.store, run, CONTEXT)['id']

    def test_offer_only_for_exact_acm_captures(self):
        self.assertIsNone(acm.propose(self.store, self.run_('s', database='Scopus'), CONTEXT))
        self.assertIsNone(acm.propose(self.store, self.run_('z', count=0), CONTEXT))
        run = self.run_('a')
        first = acm.propose(self.store, run, CONTEXT)['id']
        self.assertEqual(acm.propose(self.store, run, CONTEXT)['id'], first)
        info = acm.get(self.store, first)
        self.assertEqual((info['years'], info['types'], info['pages'], info['can_collect']),
                         ('Yıl: 2024–2026', 'Tür: research-article, short-paper', 1, True))
        self.assertEqual(acm.for_run(self.store, run['id']), first)

    def test_plan_checks_the_chrome_session_and_state(self):
        _, identity = self.offered()
        with self.assertRaises(ValidationError):
            acm.plan(self.store, identity, 'other')
        plan = acm.plan(self.store, identity, 'p')
        self.assertEqual((plan['count'], plan['query'], plan['context']['tabId']), (4, QUERY, 2))
        with self.assertRaisesRegex(ValidationError, 'toplanıyor'):
            acm.plan(self.store, identity, 'p')
        self.assertEqual(acm.fail(self.store, identity, 'HTTP 403', 'p')['state'], 'failed')
        self.assertTrue(acm.get(self.store, identity)['can_collect'])  # a failed collection may be repeated

    def test_types_are_kept_by_the_guide_and_every_record_is_listed(self):
        run, identity = self.offered(count=5)
        acm.plan(self.store, identity, 'p')
        items = [item(0), item(1, 'short-paper'), item(2, 'poster'), item(3, 'extended-abstract'), item(4, doi=item(0)['doi'])]
        job, info = acm.receive(self.store, identity, self.collection(items), 'p')
        self.assertEqual([k['doi'] for k in job['kept']], [item(0)['doi'], item(1)['doi']])
        self.assertEqual((info['state'], info['kept'], info['left_out']), ('processing', 2, 3))
        rows = list(csv.reader(io.StringIO(self.store.path(info['list_path']).read_text(encoding='utf-8-sig')), delimiter=';'))
        self.assertEqual(rows[0], acm.LIST_HEADERS)
        self.assertEqual(info['list_path'], '01_HamExport/S001_ACM_liste.csv')
        self.assertEqual([r[8] for r in rows[1:]], ['Alındı', 'Alındı', 'Alınmadı', 'Alınmadı', 'Alınmadı'])
        self.assertIn('Tür: poster', rows[3][9])
        self.assertIn('Aynı DOI', rows[5][9])

    def test_a_list_that_does_not_match_the_recorded_search_is_refused(self):
        _, identity = self.offered()
        acm.plan(self.store, identity, 'p')
        job, info = acm.receive(self.store, identity, self.collection([item(i) for i in range(3)]) | {'total': 4}, 'p')
        self.assertIsNone(job)
        self.assertEqual(info['state'], 'failed')
        self.assertIn('4 sonuç gösterdi, 3 kayıt okundu', info['message'])
        acm.plan(self.store, identity, 'p')
        job, info = acm.receive(self.store, identity, self.collection([item(i) for i in range(4)]) | {'query': 'other'}, 'p')
        self.assertIsNone(job)
        with self.assertRaises(ValidationError):
            acm.plan(self.store, identity, 'p') and acm.receive(self.store, identity, self.collection([{'doi': 'x'}] * 4), 'p')

    def test_crossref_metadata_with_list_fallback(self):
        r = acm.complete(item(0), CROSSREF)
        self.assertEqual((r['kind'], r['title'], r['date'], r['venue'], r['conference'], r['place'], r['pages'], r['abstract'], r['source']),
                         ('conferencePaper', 'Crossref Title', '2026-07-13', 'Proceedings of SIGIR', 'SIGIR 2026',
                          'Melbourne, Australia', '4045-4051', 'A multilingual collection.', 'Crossref'))
        r = acm.complete(item(0, doi='10.5555/3709347.3744082'), None)
        self.assertEqual((r['kind'], r['title'], r['date'], r['creators'][1], r['pages'], r['source']),
                         ('conferencePaper', 'Paper 0', '2026-07-19', ('Jian-Yun', 'Nie'), '4045-4051', 'ACM sonuç listesi'))
        journal = acm.complete(item(1), {'type': 'journal-article', 'container-title': ['ACM Transactions on IS'],
                                         'issued': {'date-parts': [[2025, 3]]}, 'volume': '43', 'issue': '2'})
        self.assertEqual((journal['kind'], journal['date'], journal['year']), ('journalArticle', '2025-03', '2025'))
        z = acm.zotero_item(journal)
        self.assertEqual((z['itemType'], z['publicationTitle'], z['issue'], z['DOI']), ('journalArticle', 'ACM Transactions on IS', '2', item(1)['doi']))
        self.assertNotIn('publisher', z)
        self.assertNotIn('proceedingsTitle', z)
        conference = acm.zotero_item(acm.complete(item(0), CROSSREF))
        self.assertEqual((conference['proceedingsTitle'], conference['conferenceName'], conference['ISBN']),
                         ('Proceedings of SIGIR', 'SIGIR 2026', '9798400723537'))

    def test_abstract_from_the_acm_page_comes_first_and_access_is_listed(self):
        page = dict(item(0), abstract='A multilingual collection may contain useful knowledge.', access='Açık (ACM PDF)')
        r = acm.complete(page, CROSSREF)
        self.assertEqual((r['abstract'], r['abstract_source']), ('A multilingual collection may contain useful knowledge.', 'ACM makale sayfası'))
        self.assertIn('ACM erişim: Açık (ACM PDF)', acm.zotero_item(r)['extra'])
        self.assertEqual(acm.complete(item(1), CROSSREF)['abstract_source'], 'Crossref')
        missing = acm.complete(dict(item(2), abstract_status='alınamadı (HTTP 403)'), None)
        self.assertEqual((missing['abstract'], missing['abstract_source']), ('', 'yok (alınamadı (HTTP 403))'))

    def test_progress_keeps_a_collection_alive_and_a_silent_one_can_restart(self):
        _, identity = self.offered()
        acm.plan(self.store, identity, 'p')
        acm.progress(self.store, identity, 'Özetler okunuyor: 3/26')
        info = acm.get(self.store, identity)
        self.assertEqual((info['message'], info['can_collect']), ('Özetler okunuyor: 3/26', False))
        self.store.db.execute("UPDATE acm_collections SET updated_at='2026-10-17T08:00:00+00:00' WHERE id=?", (identity,))
        self.assertTrue(acm.get(self.store, identity)['can_collect'])
        self.assertEqual(acm.plan(self.store, identity, 'p')['count'], 4)

    def test_crossref_markup_and_line_breaks_are_removed(self):
        r = acm.complete(item(0), dict(CROSSREF, title=['<scp>DoGMaTiQ:</scp>\n                    Automated Generation &amp; Nuggets'],
                                       author=[{'given': ' Bryan\n', 'family': 'Li '}]))
        self.assertEqual(r['title'], 'DoGMaTiQ: Automated Generation & Nuggets')
        self.assertEqual(r['creators'], [('Bryan', 'Li')])
        self.assertEqual(acm.plain('Cafe\u0301'), 'Caf\u00e9')

    def test_processing_writes_ris_list_zotero_and_excel_counts(self):
        run, identity = self.offered(count=3)
        acm.plan(self.store, identity, 'p')
        job, _ = acm.receive(self.store, identity, self.collection([item(0), item(1, 'short-paper'), item(2, 'poster')]), 'p')
        writer = FakeWriter()
        fetched = []
        state, message = acm.process(self.store, job, writer, fetch=lambda doi: fetched.append(doi) or (CROSSREF if doi.endswith('0') else None),
                                     sleep=lambda s: None)
        self.assertEqual(state, 'done', message)
        self.assertIn('3 kayıt okundu, 2 kayıt alındı, 1 kayıt tür dışında kaldı', message)
        self.assertIn('1 Crossref, 1 ACM listesi', message)
        self.assertEqual(fetched, [item(0)['doi'], item(1)['doi']])  # left-out records are not looked up
        ris = self.store.path('01_HamExport/S001_ACM.ris')
        self.assertEqual(ham_export.count_records(ris), 2)
        text = ris.read_text(encoding='utf-8')
        self.assertIn('TY  - CONF\nTI  - Crossref Title\nAU  - Qi, Rui\n', text)
        self.assertIn('DO  - ' + item(1)['doi'], text)
        self.assertIn('DB  - ACM Digital Library', text)
        self.assertEqual([r['doi'] for r in writer.calls[0][0]], [item(0)['doi'], item(1)['doi']])
        self.assertEqual(writer.calls[0][1], run['id'])
        rows = list(csv.reader(io.StringIO(self.store.path('01_HamExport/S001_ACM_liste.csv').read_text(encoding='utf-8-sig')), delimiter=';'))
        self.assertEqual([r[10] for r in rows[1:]], ['Crossref', 'ACM sonuç listesi', ''])
        self.assertEqual([r[11] for r in rows[1:]], ['Crossref', 'yok', ''])
        self.assertEqual(rows[0][-2:], ['Özet kaynağı', 'Erişim (ACM)'])
        old = time.time() - 10
        os.utime(ris, (old, old))
        ham_export.scan(self.store)
        info = self.store.export_info(run['id'])
        self.assertEqual((info['effective_count'], info['source']), (2, 'RIS dosyası: S001_ACM.ris'))
        self.assertEqual(info['export_files_all'], 'S001_ACM.ris; S001_ACM_liste.csv')
        self.assertEqual(acm.get(self.store, identity)['state'], 'done')
        self.assertFalse(acm.get(self.store, identity)['can_collect'])
        exported = self.store.export_excel()
        from openpyxl import load_workbook
        sheet = load_workbook(self.store.path(exported['relative_path']))['Arama_Kaydi']
        self.assertEqual(sheet.cell(2, 7).value, 'Yıl: 2024–2026\nTür: research-article, short-paper')
        self.assertEqual((sheet.cell(2, 8).value, sheet.cell(2, 9).value), (3, 2))

    def test_zotero_failure_keeps_the_files_and_allows_a_new_attempt(self):
        _, identity = self.offered(count=1)
        acm.plan(self.store, identity, 'p')
        job, _ = acm.receive(self.store, identity, self.collection([item(0)]), 'p')
        state, message = acm.process(self.store, job, FakeWriter(ValidationError('izin yok')), fetch=lambda d: None, sleep=lambda s: None)
        self.assertEqual(state, 'failed')
        self.assertIn("Zotero'ya yazılamadı: izin yok", message)
        self.assertTrue(self.store.path('01_HamExport/S001_ACM.ris').exists())
        self.assertTrue(acm.get(self.store, identity)['can_collect'])
        acm.plan(self.store, identity, 'p')
        job, _ = acm.receive(self.store, identity, self.collection([item(0)]), 'p')
        self.assertEqual(acm.process(self.store, job, FakeWriter(), fetch=lambda d: None, sleep=lambda s: None)[0], 'done')
        self.assertEqual(sorted(p.name for p in self.store.path('01_HamExport').iterdir()), ['S001_ACM.ris', 'S001_ACM_liste.csv'])

    def test_zotero_writer_puts_acm_records_in_their_own_search_collection(self):
        run, identity = self.offered(count=2)
        api = FakeLocalAPI()
        credentials = {}
        for p in (patch('literature.zotero_writer.load_api_key', side_effect=lambda name: credentials.get(name, '')),
                  patch('literature.zotero_writer.save_api_key', side_effect=lambda key, name: credentials.__setitem__(name, key))):
            p.start()
            self.addCleanup(p.stop)
        writer = ZoteroWriter(self.store, api)
        self.store.update_settings({'zotero_mode': 'local', 'library': 'groups/6718452',
                                    'collection_mapping': {'candidate': 'BX38VWI6'}})
        writer.preflight(authorize=True)
        records = [{'doi': r['doi'], 'item': acm.zotero_item(acm.complete(r, None))} for r in (item(0), item(1, 'short-paper'))]
        result = writer.import_acm(records, run['id'])
        self.assertEqual((result['imported'], result['failed']), (2, 0), result['failures'])
        collection = next(c for c in api.objects['collections'].values() if c['data']['name'] == 'S001_ACM')
        self.assertEqual(collection['data']['parentCollection'], 'U8THQFWA')
        items = [o['data'] for o in api.objects['items'].values()]
        self.assertEqual(sorted(i['DOI'] for i in items), sorted(r['doi'] for r in records))
        self.assertTrue(all(set(i['collections']) == {'BX38VWI6', collection['key']} for i in items))
        again = writer.import_acm(records, run['id'])
        self.assertEqual((again['imported'], again['reused']), (0, 2))
        # Zotero normalizing a text field must not turn a written item into a "conflict" (Deneme, 17.10.2026)
        original = api.seed
        api.seed = lambda kind, data: original(kind, dict(data, title=' '.join(data['title'].split()) + ' ') if kind == 'items' else data)
        extra = [{'doi': r['doi'], 'item': acm.zotero_item(acm.complete(r, None))} for r in (item(5), item(6))]
        result = writer.import_acm(records + extra, run['id'])
        self.assertEqual((result['imported'], result['reused'], result['failed']), (2, 2, 0), result['failures'])
        # A lost reply after Zotero committed: the next collection finds the items instead of copying them
        api.seed = original
        api.timeout_after_items = True
        lost = [{'doi': r['doi'], 'item': acm.zotero_item(acm.complete(r, None))} for r in (item(7), item(8))]
        self.assertEqual(writer.import_acm(records + extra + lost, run['id'])['failed'], 2)
        healed = writer.import_acm(records + extra + lost, run['id'])
        self.assertEqual((healed['imported'], healed['reused'], healed['failed']), (0, 6, 0), healed['failures'])
        self.assertEqual(len(api.objects['items']), 6)
        with self.assertRaises(ValidationError):
            writer.import_acm(records, self.run_('s', database='Scopus')['id'])


if __name__ == '__main__':
    unittest.main()
