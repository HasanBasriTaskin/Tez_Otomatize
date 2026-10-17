import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook
from literature.storage import Store, ValidationError, SEARCH_HEADERS


def fields(**updates):
    result = dict(database='Scopus', query_raw='TITLE-ABS-KEY("Türkçe" AND RAG)',
                  query_verified=True, fields='TITLE-ABS-KEY', filters=['2022–2026'],
                  count_raw='0 results', count_value=0, count_qualifier='exact',
                  safe_url='https://www.scopus.com/results', status='complete')
    result.update(updates)
    return result


class StorageTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = Store(self.root)
        self.png = self.root / 'input.png'
        Image.new('RGB', (320, 200), 'white').save(self.png)

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def capture(self, job='job-1', context=None):
        context = context or dict(profile='p', windowId=1, tabId=2, documentId='d', epoch=0)
        self.store.begin_job(job, context)
        self.store.save_capture(job, self.png.read_bytes(), captured_at='2026-10-14T21:05:00+00:00')
        return self.store.complete_job(job, fields(), context)

    def test_zero_and_local_capture_date(self):
        result = self.capture()
        self.assertEqual(result['display_id'], 'S001')
        self.assertEqual(result['count_value'], 0)
        cap = self.store.captures(result['id'])[0]
        self.assertIn('S001_20261015_a_sonuc.png', cap['relative_path'])
        self.assertEqual(cap['width'], 320)
        self.assertEqual(len(cap['sha256']), 64)

    def test_retry_is_idempotent_and_extra_capture_keeps_one_run(self):
        first = self.capture()
        self.capture()
        second = self.capture('job-2')
        self.assertEqual(first['id'], second['id'])
        self.assertEqual(len(self.store.list_runs()), 1)
        caps = self.store.captures(first['id'])
        self.assertEqual(len(caps), 2)
        self.assertNotEqual(caps[0]['relative_path'], caps[1]['relative_path'])
        self.assertTrue(any('_v02' in c['relative_path'] for c in caps))

    def test_changed_tab_keeps_image_pending(self):
        a = dict(profile='p', windowId=1, tabId=2, documentId='d', epoch=0)
        b = dict(a, tabId=3)
        self.store.begin_job('changed', a)
        self.store.save_capture('changed', self.png.read_bytes())
        r = self.store.complete_job('changed', fields(), b)
        self.assertEqual(r['state'], 'incomplete')
        self.assertEqual(len(self.store.list_runs()), 0)
        self.assertTrue(self.store.path(r['image_path']).exists())

    def test_unknown_count_does_not_become_zero(self):
        self.store.begin_job('null', {})
        self.store.save_capture('null', self.png.read_bytes())
        r = self.store.complete_job('null', fields(count_value=None), {})
        self.assertEqual(r['state'], 'incomplete')

    def test_modified_pending_png_cannot_be_verified(self):
        self.store.begin_job('modified', {})
        job=self.store.save_capture('modified', self.png.read_bytes())
        self.store.path(job['image_path']).write_bytes(b'changed')
        result=self.store.complete_job('modified', fields(), {})
        self.assertEqual(result['state'], 'incomplete')
        self.assertEqual(len(self.store.list_runs()), 0)

    def test_different_query_and_new_day_are_new_runs(self):
        first = self.capture()
        c = dict(profile='p', windowId=1, tabId=2, documentId='d', epoch=0)
        self.store.begin_job('new-day', c)
        self.store.save_capture('new-day', self.png.read_bytes(), captured_at='2026-10-16T12:00:00+00:00')
        second = self.store.complete_job('new-day', fields(), c)
        self.assertNotEqual(first['id'], second['id'])

    def test_excel_preserves_text_and_portable_links(self):
        self.capture()
        self.store.update_run('S001', {'notes': '=HYPERLINK("bad")'})
        out = self.store.export_excel()
        w = load_workbook(out['path'])
        self.assertEqual(w.sheetnames, ['Arama_Kaydi', 'Matris', 'Hariç'])
        s = w['Arama_Kaydi']
        self.assertEqual([c.value for c in s[1]], SEARCH_HEADERS)
        self.assertEqual(s['A2'].value, 'S001')
        self.assertEqual(s['H2'].value, 0)
        self.assertEqual(s['K2'].data_type, 's')
        self.assertTrue(s['J2'].hyperlink.target.startswith('../02_Ekranlar/'))
        w.close()

    def test_manual_export_count_operator_and_files_reach_excel(self):
        self.store.update_settings({'operator': 'Hasan Basri Taşkın'})
        self.capture()
        self.store.update_run('S001', {'manual_count': '245', 'export_files': 'S001_Scopus_001-245.ris'})
        info = self.store.export_info(self.store.get_run('S001')['id'])
        self.assertEqual((info['effective_count'], info['source']), (245, 'Elle girildi'))
        w = load_workbook(self.store.export_excel()['path'])
        try:
            s = w['Arama_Kaydi']
            self.assertEqual(s.max_column, len(SEARCH_HEADERS))
            self.assertEqual(s['I2'].value, 245)
            self.assertEqual(s['L2'].value, 'Elle girildi')
            self.assertEqual(s['M2'].value, 'S001_Scopus_001-245.ris')
            self.assertEqual(s['N2'].value, 'Hasan Basri Taşkın')
        finally:
            w.close()

    def test_zotero_count_used_when_no_manual_count(self):
        self.capture()
        run = self.store.get_run('S001')
        with self.store.db:
            self.store.db.execute("INSERT INTO search_run_exports(run_id,zotero_count,zotero_collection) VALUES(?,?,?)", (run['id'], 248, 'S001'))
        info = self.store.export_info(run['id'])
        self.assertEqual(info['effective_count'], 248)
        self.assertIn('Zotero', info['source'])
        self.store.update_run('S001', {'manual_count': 250})
        info = self.store.export_info(run['id'])
        self.assertEqual(info['effective_count'], 250)
        self.assertIn('farkı kontrol edin', info['source'])
        self.store.update_run('S001', {'manual_count': ''})
        self.assertEqual(self.store.export_info(run['id'])['effective_count'], 248)

    def test_publication_listed_twice_in_one_ris_is_explained_not_flagged(self):
        from literature import zotero_ris
        self.capture()
        run = self.store.get_run('S001')
        zotero_ris.ensure_schema(self.store)
        with self.store.db:
            self.store.db.execute("INSERT INTO ham_export_files VALUES(?,?,?,?,?,?,?)", ('01_HamExport/S001_Scopus.ris', run['id'], 849, 0, 1, 1, 'x'))
            self.store.db.execute("INSERT INTO search_run_exports(run_id,zotero_count,zotero_collection) VALUES(?,?,?)", (run['id'], 848, 'S001_Scopus'))
            self.store.db.execute("INSERT INTO zotero_ris_imports(run_id,display_id,files,records,unique_items,updated_at) VALUES(?,?,?,?,?,?)",
                                  (run['id'], 'S001', '[]', 849, 848, 'x'))
        source = self.store.export_info(run['id'])['source']
        self.assertIn("dosyada 1 kayıt iki kez geçtiği için Zotero'da 848 öğe", source)
        self.assertNotIn('farkı kontrol edin (Zotero', source.replace('ekranda', ''))
        self.assertNotIn('Zotero 848 kayıt; farkı kontrol edin', source)
        with self.store.db:
            self.store.db.execute('UPDATE zotero_ris_imports SET unique_items=849')
        self.assertIn('Zotero 848 kayıt; farkı kontrol edin', self.store.export_info(run['id'])['source'])

    def test_workbook_properties_carry_no_tool_description(self):
        self.capture()
        w = load_workbook(self.store.export_excel()['path'])
        try:
            self.assertFalse(w.properties.description)
        finally:
            w.close()

    def test_evidence_fields_and_bad_counts_rejected(self):
        self.capture()
        for bad in [{'query_raw': 'x'}, {'count_value': 5}, {'manual_count': -1}, {'manual_count': 'çok'},
                    {'export_files': 'x' * 2001}, {}]:
            with self.assertRaises(ValidationError):
                self.store.update_run('S001', bad)

    def test_excel_user_edits_are_not_overwritten(self):
        self.capture()
        p = Path(self.store.export_excel()['path'])
        p.write_bytes(b'user changed file')
        r = self.store.export_excel()
        self.assertNotEqual(str(p), r['path'])
        self.assertEqual(p.read_bytes(), b'user changed file')

    def test_excel_project_metadata_stays_literal(self):
        self.store.update_settings({'project_name':'=1+1','operator':'=HYPERLINK("bad")'})
        out=self.store.export_excel()
        wb=load_workbook(out['path'])
        try:
            self.assertEqual(wb.properties.title, '=1+1')
            self.assertEqual(wb.properties.creator, '=HYPERLINK("bad")')
            for sheet in wb:
                self.assertEqual(sheet.max_row, 1)
                self.assertFalse(any(c.data_type == 'f' for row in sheet for c in row))
        finally:wb.close()

    def test_excel_team_change_uses_new_filename(self):
        before=self.store.export_excel()
        self.store.update_settings({'team_code':'NEW-TEAM'})
        after=self.store.export_excel()
        self.assertIn('NEW-TEAM',Path(after['path']).name)
        self.assertTrue(Path(before['path']).exists())

    def test_disk_write_failure_preserves_incomplete_job(self):
        self.store.begin_job('disk-full',{})
        with patch.object(self.store,'_atomic_write',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):self.store.save_capture('disk-full',self.png.read_bytes())
        self.assertEqual(self.store.get_job('disk-full')['state'],'incomplete')
        self.assertEqual(self.store.list_runs(),[])

    def test_locked_excel_creates_separate_version(self):
        self.capture()
        first=self.store.export_excel();old=Path(first['path']).read_bytes()
        original=self.store._atomic_write
        def locked(relative,data,replace=False):
            if relative==first['relative_path']:raise PermissionError('Excel lock')
            return original(relative,data,replace)
        with patch.object(self.store,'_atomic_write',side_effect=locked):
            new=self.store.export_excel()
        self.assertEqual(new['status'],'locked_versioned')
        self.assertNotEqual(first['path'],new['path'])
        self.assertEqual(Path(first['path']).read_bytes(),old)

    def test_open_excel_owner_file_creates_separate_version(self):
        self.capture()
        first = Path(self.store.export_excel()['path'])
        old = first.read_bytes()
        first.with_name('~$' + first.name).write_bytes(b'Excel owner')
        result = self.store.export_excel()
        self.assertEqual(result['status'], 'locked_versioned')
        self.assertNotEqual(str(first), result['path'])
        self.assertEqual(first.read_bytes(), old)

    def test_backup_restores_relative_excel_and_image_links(self):
        from literature.backup import create_backup,restore_backup
        self.capture();excel=self.store.export_excel()
        archive=create_backup(self.store)
        with tempfile.TemporaryDirectory() as other:
            target=Path(other)/'restored';restore_backup(archive,target)
            wb=load_workbook(target/excel['relative_path'])
            try:
                gallery=((target/excel['relative_path']).parent/wb['Arama_Kaydi']['J2'].hyperlink.target).resolve()
                self.assertTrue(gallery.is_file())
                self.assertTrue((gallery.parent/'S001_20261015_a_sonuc.png').is_file())
            finally:wb.close()

    def test_backup_skips_excel_owner_lock_file(self):
        import zipfile
        from literature.backup import create_backup
        self.capture()
        excel = Path(self.store.export_excel()['path'])
        owner = excel.with_name('~$' + excel.name)
        owner.write_bytes(b'temporary Excel owner lock')
        read_bytes = Path.read_bytes
        def locked(path):
            if path == owner:
                raise PermissionError('Excel holds its owner file exclusively')
            return read_bytes(path)
        with patch.object(Path, 'read_bytes', locked):
            archive = create_backup(self.store)
        with zipfile.ZipFile(archive) as backup:
            self.assertNotIn(owner.relative_to(self.root).as_posix(), backup.namelist())
            self.assertIn(excel.relative_to(self.root).as_posix(), backup.namelist())

    def test_path_traversal_rejected(self):
        for name in ['../out.png', 'C:/out.png', '/out.png']:
            with self.assertRaises(ValidationError):
                self.store.path(name)

    def test_recovery_after_rename_before_database_commit(self):
        self.store.begin_job('recover', {})
        self.store.save_capture('recover', self.png.read_bytes())
        with patch.object(self.store, '_finish_commit', side_effect=RuntimeError('power loss')):
            with self.assertRaises(RuntimeError):
                self.store.complete_job('recover', fields(), {})
        self.store.close()
        self.store = Store(self.root)
        self.assertEqual(len(self.store.list_runs()), 1)
        self.assertEqual(len(self.store.captures(self.store.list_runs()[0]['id'])), 1)
        self.assertEqual(self.store.get_job('recover')['state'], 'verified')


if __name__ == '__main__':
    unittest.main()
