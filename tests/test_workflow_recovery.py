import tempfile
import unittest
from unittest.mock import patch

from literature.capture import CaptureController
from literature.storage import Store, ValidationError, json_text
from literature.workflow import WorkflowEngine, ClockGuard
from test_capture import Desktop, Peer
from test_storage import fields


class WorkflowPeer(Peer):
    connected=True
    query='';stage='search'
    def request(self,kind,payload,timeout=10):
        if kind=='workflow_guard':return {'ready':True}
        if kind=='workflow_browser':
            action=payload['action']
            if action=='fill_search':self.query=payload['query']
            if action=='submit_search':self.stage='results'
            if action=='inspect':return {'stage':self.stage,'data':fields(query_raw=self.query,count_value=0)}
            if action=='prepare_evidence':return {'ready':True,'evidence':{'requiresExtraViewport':False}}
            return {'stage':self.stage}
        result=super().request(kind,payload,timeout)
        if kind=='extract_capture':result['data']=fields(query_raw=self.query,count_value=0)
        return result


class RecoveryTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.store=Store(self.tmp.name)
        self.desktop=Desktop();self.peer=WorkflowPeer(self.desktop)
        self.capture=CaptureController(self.store,self.desktop,lambda *args:None)
        self.capture.bind(100,7,self.peer)
        self.engine=WorkflowEngine(self.store,self.capture);self.engine.stable_seconds=0
        self.session=self.engine.records.create({},dict(profile='p',windowId=7,tabId=9))
        self.engine.records.set_session(self.session['id'],state='running')
        self.engine.identity=self.session['id'];self.engine.peer=self.peer;self.engine.hwnd=100
        self.engine.epoch=self.desktop.epoch;self.engine.clock=ClockGuard()

    def tearDown(self):self.store.close();self.tmp.cleanup()

    def test_excel_failure_on_final_step_cannot_mark_session_completed(self):
        records=self.engine.records
        records.advance(self.session['id'],0,'P00')
        for i in range(1,6):records.advance(self.session['id'],0,'Q0'+str(i))
        with patch.object(self.store,'export_excel',side_effect=PermissionError('disk full')):
            self.engine.run(self.session['id'],self.peer)
        current=records.get(self.session['id'])
        self.assertEqual(current['state'],'paused')
        self.assertNotEqual(current['steps'][-1]['state'],'done')
        self.assertEqual(len(self.store.list_runs()),1)
        records.set_session(current['id'],state='running')
        self.engine.run(current['id'],self.peer)
        self.assertEqual(records.get(current['id'])['state'],'completed')
        self.assertEqual(len(self.store.list_runs()),1,'Resume duplicated the verified capture')

    def test_excel_failure_stops_before_export_or_zotero_write(self):
        self.engine.records.advance(self.session['id'],0,'P00')
        data=fields(count_value=1)
        with patch.object(self.engine,'_search',return_value=data), patch.object(self.engine,'_capture',return_value={'id':'run','count_value':1}), patch.object(self.store,'export_excel',side_effect=PermissionError('disk full')), patch.object(self.engine,'_export') as export:
            self.engine.run(self.session['id'],self.peer)
        export.assert_not_called()
        self.assertEqual(self.engine.records.get(self.session['id'])['state'],'paused')

    def test_verified_capture_is_relinked_after_restart_before_step_save(self):
        step=self.session['steps'][0];self.peer.query=step['query']
        run=self.capture.capture(automatic={'job_id':'saved-before-crash','context':self.session['context'],'query':step['query']})
        self.engine.records.update_step(step['id'],capture_job_id='saved-before-crash',state='capturing')
        step=self.engine.records.get(self.session['id'])['steps'][0]
        self.engine._capture(step,fields(query_raw=step['query'],count_value=0))
        restored=self.engine.records.get(self.session['id'])['steps'][0]
        self.assertEqual(restored['run_id'],run['id'])
        self.assertTrue(restored['detail']['evidence_complete'])
        self.assertEqual(len(self.store.list_runs()),1)

    def test_unfinished_extra_viewport_blocks_progress_and_recovers_same_run(self):
        step=self.session['steps'][0];self.peer.query=step['query']
        run=self.capture.capture(automatic={'job_id':'first-viewport','context':self.session['context'],'query':step['query']})
        self.engine.records.update_step(step['id'],capture_job_id='first-viewport',run_id=run['id'],state='capturing',detail_json=json_text({'extra_evidence_required':True}))
        step=self.engine.records.get(self.session['id'])['steps'][0]
        with patch.object(self.engine,'_prepare_evidence',side_effect=ValidationError('count not visible')):
            with self.assertRaises(ValidationError):self.engine._capture(step,fields(query_raw=step['query'],count_value=0))
        self.assertNotIn('evidence_complete',self.engine.records.get(self.session['id'])['steps'][0]['detail'])
        self.engine._capture(step,fields(query_raw=step['query'],count_value=0))
        self.assertEqual(len(self.store.list_runs()),1)
        self.assertEqual(len(self.store.captures(run['id'])),2)

    def test_focus_change_even_when_returned_pauses(self):
        self.desktop.epoch+=1
        with self.assertRaisesRegex(ValidationError,'pencere'):self.engine.check()

    def test_switching_tabs_in_same_window_stops_before_next_operation(self):
        request=self.peer.request
        def changed(kind,payload,timeout=10):
            if kind=='workflow_guard':return {'ready':False}
            return request(kind,payload,timeout)
        self.peer.request=changed
        with self.assertRaisesRegex(ValidationError,'sekmesi'):self.engine.check()

    def test_excel_exported_count_is_actual_verified_ris_sum(self):
        from openpyxl import load_workbook
        step=self.session['steps'][0];self.peer.query=step['query']
        run=self.capture.capture(automatic={'job_id':'excel-count','context':self.session['context'],'query':step['query']})
        self.engine.records.update_step(step['id'],run_id=run['id'])
        with self.store.db:
            self.store.db.execute("INSERT INTO workflow_exports(id,step_id,range_start,range_end,state,record_count) VALUES('e',?,1,3,'complete',3)",(step['id'],))
        output=self.store.export_excel();wb=load_workbook(output['path'])
        try:self.assertEqual(wb['Arama_Kaydi'].cell(2,9).value,3);self.assertEqual(wb['Arama_Kaydi'].max_column,14);self.assertEqual(wb['Arama_Kaydi'].cell(2,12).value,'Otomatik RIS doğrulaması')
        finally:wb.close()
