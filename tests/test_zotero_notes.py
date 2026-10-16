import tempfile
import unittest
from pathlib import Path

from PIL import Image
from literature.storage import Store
from literature.zotero import (EXCLUSION_TEMPLATE, READING_TEMPLATE, UNKNOWN_KEY, import_items,
                               parse_note, search_codes)


def note(parent, key, text):
    html = ''.join(f'<p>{line}</p>' for line in text.splitlines())
    return {'key': key, 'data': {'itemType': 'note', 'parentItem': parent, 'note': html}}


def paper(key, collections, title='Türkçe RAG'):
    return {'key': key, 'data': {'itemType': 'conferencePaper', 'title': title, 'date': '2025',
                                 'creators': [{'creatorType': 'author', 'firstName': 'Ayşe', 'lastName': 'Köse'}],
                                 'collections': collections}}


class NoteParsingTest(unittest.TestCase):
    def test_untouched_templates_parse_as_empty_fields(self):
        reading = parse_note(READING_TEMPLATE)
        self.assertEqual(reading['Matris ID'], '')
        self.assertEqual(reading['Karar'], 'Dahil')
        self.assertEqual(reading['Okuma düzeyi'], '')
        self.assertNotIn(UNKNOWN_KEY, reading)
        exclusion = parse_note(EXCLUSION_TEMPLATE)
        self.assertEqual((exclusion['Karar'], exclusion['Aşama'], exclusion['Gerekçe kodu']), ('Hariç', 'Tam metin', ''))

    def test_old_notes_still_read_and_aciklama_alias(self):
        old = parse_note('<p>Amaç/problem: Deneme</p><p>Dahil/hariç gerekçesi: H2 deney yok</p>')
        self.assertEqual(old['Dahil/hariç gerekçesi'], 'H2 deney yok')
        new = parse_note('<p>Açıklama: Yalnız sistem tanıtımı</p>')
        self.assertEqual(new['Dahil/hariç gerekçesi'], 'Yalnız sistem tanıtımı')

    def test_unknown_field_kept_with_previous_and_reported(self):
        parsed = parse_note('<p>Temel bulgu: Recall@5 = 0,81</p><p>Tablo 3: s. 6</p><p>Gerekce: yazım hatası</p>')
        self.assertIn('Tablo 3: s. 6', parsed['Temel bulgu'])
        self.assertIn('Gerekce: yazım hatası', parsed['Temel bulgu'])
        self.assertEqual(parsed[UNKNOWN_KEY], 'Gerekce')

    def test_search_codes_from_collection_names(self):
        self.assertEqual(search_codes(['S004', 'S012_Scopus', '01_Aday', 'S1', None, 'S0045x']), ['S004', 'S0045', 'S012'])


class ZoteroImportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = Store(self.root)
        png = self.root / 'x.png'
        Image.new('RGB', (40, 30), 'white').save(png)
        for job in ('a', 'b'):
            self.store.begin_job(job, {})
            self.store.save_capture(job, png.read_bytes(), captured_at='2026-10-16T10:00:00+00:00')
            self.store.complete_job(job, dict(database='Scopus', query_raw='TITLE-ABS-KEY(x)' + job, query_verified=True,
                                              fields='TITLE-ABS-KEY', filters=[], count_raw='2', count_value=2,
                                              count_qualifier='exact', status='complete'), {}, new_run=True)
        self.collections = {'C1': 'S001', 'C2': 'S002_IEEE', 'C9': 'S099', 'IN': '03_TamMetin_Dahil',
                            'EX': '04_Haric_TamMetin', 'ET': '05_Haric_BaslikOzet', 'TA': '02_BaslikOzet_Gecen'}
        self.mapping = {'included': 'IN', 'excluded': 'EX', 'excluded_title': 'ET', 'title_abstract': 'TA'}

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def state(self, key):
        return self.store.db.execute('SELECT state FROM papers WHERE item_key=?', (key,)).fetchone()[0]

    def test_roles_codes_counts_and_attachments(self):
        items = [paper('P1', ['C1', 'C2', 'IN']), paper('P2', ['C1', 'ET']), paper('P3', ['C9', 'TA']),
                 paper('P4', ['C1', 'IN', 'ET']), paper('P5', ['C2', 'EX', 'TA']),
                 note('P5', 'N5', EXCLUSION_TEMPLATE.replace('(H1 / H2 / H3 / H4 / H5 / H6)', 'H2')),
                 {'key': 'A1', 'data': {'itemType': 'attachment', 'parentItem': 'P1', 'contentType': 'application/pdf',
                                        'linkMode': 'imported_url', 'filename': 'Kose2025.pdf'}},
                 {'key': 'A2', 'data': {'itemType': 'attachment', 'parentItem': 'P1', 'contentType': 'text/html',
                                        'linkMode': 'imported_url', 'filename': 'snapshot.html'}}]
        result = import_items(self.store, 'groups/1', items, self.mapping, self.collections)
        self.assertEqual(result['unknown_codes'], ['S099'])
        self.assertEqual([self.state(k) for k in ('P1', 'P2', 'P3', 'P4', 'P5')],
                         ['included', 'excluded_title', 'title_passed', 'conflict', 'excluded'])
        s1 = self.store.export_info(self.store.get_run('S001')['id'])
        s2 = self.store.export_info(self.store.get_run('S002')['id'])
        self.assertEqual((s1['zotero_count'], s2['zotero_count']), (3, 2))
        self.assertEqual(s2['zotero_collection'], 'S002_IEEE')
        codes = [r[0] for r in self.store.db.execute(
            "SELECT code FROM paper_search_codes c JOIN papers p ON p.id=c.paper_id WHERE p.item_key='P1' ORDER BY code")]
        self.assertEqual(codes, ['S001', 'S002'])
        pdfs = self.store.db.execute(
            "SELECT attachment_key,filename FROM paper_attachments a JOIN papers p ON p.id=a.paper_id WHERE p.item_key='P1'").fetchall()
        self.assertEqual([tuple(r) for r in pdfs], [('A1', 'Kose2025.pdf')])
        # A later read without the S collection clears the stale Zotero count.
        import_items(self.store, 'groups/1', [paper('P1', ['IN'])], self.mapping, self.collections)
        self.assertIsNone(self.store.export_info(self.store.get_run('S001')['id'])['zotero_count'])

    def test_items_of_a_previous_library_stop_counting(self):
        # 16.10.2026: "Zotero'dan oku" ran once on users/0 before the TEZ group was set.
        import_items(self.store, 'users/0', [paper('MINE', [])], {}, {})
        self.assertEqual(self.state('MINE'), 'candidate')
        import_items(self.store, 'groups/1', [paper('P1', ['IN'])], self.mapping, self.collections)
        self.assertEqual((self.state('MINE'), self.state('P1')), ('unavailable', 'included'))
        states = dict(self.store.db.execute('SELECT state,COUNT(*) FROM papers GROUP BY state').fetchall())
        self.assertNotIn('candidate', states)

    def test_empty_duplicate_template_field_is_not_a_conflict(self):
        filled = READING_TEMPLATE.replace('(L001, L002 … dahil edince verin)', 'L007')
        items = [paper('P1', ['IN']), note('P1', 'N1', filled), note('P1', 'N2', READING_TEMPLATE)]
        import_items(self.store, 'groups/1', items, self.mapping, self.collections)
        self.assertEqual(self.state('P1'), 'included')


    def test_excel_matrix_and_exclusion_sheets(self):
        from openpyxl import load_workbook
        from literature.storage import EXCLUDED_HEADERS, MATRIX_HEADERS
        reading = (READING_TEMPLATE.replace('(L001, L002 … dahil edince verin)', 'L002')
                   .replace('(Dergi makalesi / Konferans bildirisi / Ön baskı / Derleme)', 'Derleme')
                   .replace('Okuyan kişi: ', 'Okuyan kişi: Esra Sultan Doğan'))
        older = '\n'.join(['Dahil/hariç gerekçesi: H5 yıl dışı', 'Okuma düzeyi: 1'])
        exclusion = (EXCLUSION_TEMPLATE.replace('(H1 / H2 / H3 / H4 / H5 / H6)', 'H2')
                     .replace('Açıklama: ', 'Açıklama: Deney ve ölçüt yok').replace('Karar veren: ', 'Karar veren: Hasan Basri Taşkın')
                     .replace('(GG.AA.YYYY)', '20.10.2026'))
        items = [paper('P1', ['C1', 'IN'], 'RAGAs: Automated Evaluation'), note('P1', 'N1', reading),
                 paper('P0', ['C2', 'IN'], 'Unnumbered paper'),
                 paper('P2', ['C2', 'EX']), note('P2', 'N2', exclusion),
                 paper('P3', ['C1', 'EX']), note('P3', 'N3', older),
                 paper('P4', ['C1', 'ET'])]
        import_items(self.store, 'groups/1', items, self.mapping, self.collections)
        wb = load_workbook(self.store.export_excel()['path'])
        try:
            m, h = wb['Matris'], wb['Hariç']
            self.assertEqual([c.value for c in m[1]], MATRIX_HEADERS)
            self.assertEqual([c.value for c in h[1]], EXCLUDED_HEADERS)
            self.assertEqual([m.cell(r, 1).value for r in (2, 3)], ['L002', 'P0'])
            row = {MATRIX_HEADERS[i]: m.cell(2, i + 1).value for i in range(len(MATRIX_HEADERS))}
            self.assertEqual((row['Yıl'], row['Yayın türü'], row['Okuyan üye']), ('2025', 'Derleme', 'Esra Sultan Doğan'))
            self.assertEqual(row['Bulunma yolu / arama kaydı'], 'S001')
            self.assertEqual(row['Yerel PDF dosyası'], 'L002_Kose_2025_RAGAs.pdf')
            self.assertEqual(m.cell(3, 5).value, 'Konferans bildirisi')
            self.assertEqual(h.max_row, 3)  # title/abstract exclusions are not listed
            rows = {h.cell(r, 1).value: [h.cell(r, c).value for c in range(1, 10)] for r in (2, 3)}
            self.assertEqual(rows['P2'][3:], ['S002', 'Tam metin', 'H2', 'Deney ve ölçüt yok', 'Hasan Basri Taşkın', '20.10.2026'])
            self.assertEqual(rows['P3'][4:6], ['Tam metin', 'H5'])  # old note: code read from the explanation
        finally:
            wb.close()

if __name__ == '__main__':
    unittest.main()
