import io
import json
import tempfile
import unittest

from openpyxl import load_workbook

from literature import screening, screening_workbook
from literature.storage import Store
from tests.test_screening_apply import decisions


class ScreeningWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(self.tmp.name)
        self.path = self.store.path(screening.RELATIVE.as_posix())
        self.path.parent.mkdir(parents=True)
        data = decisions()
        data['kayitlar'][1].update(no='12', arama_konusu='Q01 Türkçe', AS='AS1;AS2', secim_gerekcesi='Türkçe parçalama',
                                   notlar='Kısa not', tur='journalArticle', yil='2025')
        self.path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        self.path.with_name('eleme_olcutleri.md').write_text('# Ölçütler\n- D1: RAG', encoding='utf-8')

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_sheets_and_counts_follow_the_decisions(self):
        result = screening_workbook.export(self.store)
        wb = load_workbook(result['path'])
        self.assertEqual(wb.sheetnames, ['Ozet', 'TamMetin_1', 'Yedek', 'Haric_BaslikOzet', 'Olcutler'])
        summary = {row[0]: row[1] for row in wb['Ozet'].iter_rows(values_only=True) if row[0]}
        self.assertEqual((summary['Başlıktan hariç tutulan'], summary['Özetten hariç tutulan'], summary['   Yedek']), (1, 1, 1))
        self.assertEqual(summary['Kontrol: hariç + geçen + özeti okunmayan = aday'], 'Doğru')
        self.assertTrue(summary['Onay'].startswith('Onaylı'))
        selected = list(wb['TamMetin_1'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(selected[0][:3], (12, 'BBBB2222', 'Chunking for Turkish QA'))
        self.assertEqual(selected[0][8], 'Türkçe parçalama')
        self.assertEqual(wb['Haric_BaslikOzet'].max_row, 3)
        self.assertEqual(wb['Olcutler']['A1'].value, 'Ölçütler')

    def test_unread_records_get_their_own_sheet(self):
        data = decisions(approved=False)
        data['kayitlar'].append(dict(key='EEEE2222', baslik='Unread', durum='Özet okunmadı'))
        wb = load_workbook(io.BytesIO(screening_workbook.build(data)))
        self.assertIn('Ozet_Okunmadi', wb.sheetnames)
        self.assertNotIn('Olcutler', wb.sheetnames)

    def test_a_file_the_program_did_not_write_is_kept(self):
        target = self.store.path(screening_workbook.RELATIVE)
        target.write_bytes(b'user file')
        result = screening_workbook.export(self.store)
        self.assertEqual(result['status'], 'versioned')
        self.assertEqual(target.read_bytes(), b'user file')
        self.path.unlink()
        target.unlink()
        self.path.write_text(json.dumps(decisions(), ensure_ascii=False), encoding='utf-8')
        first = screening_workbook.export(self.store)
        again = screening_workbook.export(self.store)
        self.assertEqual((first['status'], again['status'], again['path']), ('updated', 'updated', str(target)))


if __name__ == '__main__':
    unittest.main()
