import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from literature.storage import Store
from literature.zotero import import_items


def items():
    note = ('<p>Matris ID: L001</p><p>Karar: Dahil</p><p>Yayın türü: Konferans bildirisi</p>'
            '<p>Kalite sınıfı ve yılı: CORE A* (2023)</p><p>Kalite kaynağı: https://portal.core.edu.au</p>'
            '<p>Okuyan kişi: Hasan Basri Taşkın</p><p>Okuma düzeyi: 3</p><p>Okuma tarihi: 14.10.2026</p>'
            '<p>Çekirdek adayı: Evet</p><p>Amaç/problem: Türkçe soru-cevap.</p><p>Temel bulgu: MRR = 0,38 (Tablo 2, s. 5)</p>'
            '<p>Kanıt sayfaları: s. 5</p><p>Yerel PDF: L001_Aksoy_2026_Benchmark.pdf</p>')
    return [{'key': 'ABCD2345', 'data': {'itemType': 'conferencePaper', 'title': 'Benchmark RAG', 'date': '2026',
                                         'creators': [{'creatorType': 'author', 'firstName': 'M.', 'lastName': 'Aksoy'}],
                                         'url': 'https://www.scopus.com/x', 'collections': ['IN']},
             'bib': '<div>[1]</div><div>M. Aksoy, “Benchmark RAG,” 2026.</div>'},
            {'key': 'NOTE2345', 'data': {'itemType': 'note', 'parentItem': 'ABCD2345', 'note': note}},
            {'key': 'EXCL2345', 'data': {'itemType': 'conferencePaper', 'title': 'Closed paper', 'date': '2025',
                                         'DOI': '10.1007/x', 'collections': ['EX']}},
            {'key': 'NOTX2345', 'data': {'itemType': 'note', 'parentItem': 'EXCL2345', 'note':
                '<p>Karar: Hariç</p><p>Aşama: Tam metin</p><p>Gerekçe kodu: H3</p><p>Açıklama: erişilemedi</p>'
                '<p>Karar veren: Hasan Basri Taşkın</p><p>Karar tarihi: 17.10.2026</p>'}}]


class DeliveryWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.tmp.name) / 'p')
        import_items(self.store, 'users/1', items(), {'included': 'IN', 'excluded': 'EX'})

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_only_the_guide_columns_and_overwrite(self):
        result = self.store.export_delivery_workbook()
        self.assertEqual(result['relative_path'], '05_Teslim/OEC-1017_O2_LiteraturMatrisi.xlsx')
        wb = load_workbook(result['path'])
        try:
            self.assertEqual(wb.sheetnames, ['Arama_Kaydi', 'Matris', 'Hariç'])
            self.assertEqual([c.value for c in wb['Arama_Kaydi'][1]],
                             ['No', 'Tarih', 'Kaynak', 'Sorgu / işlem', 'Alan ve filtreler', 'Sonuç', 'Aktarılan'])
            matrix = wb['Matris']
            headers = [c.value for c in matrix[1]]
            self.assertEqual(len(headers), 20)
            for absent in ('Okuma tarihi', 'Kalite kaynağı', 'Çekirdek adayı', 'İstisna gerekçesi'):
                self.assertNotIn(absent, headers)
            row = {h: c.value for h, c in zip(headers, matrix[2])}
            self.assertEqual(row['ID'], 'L001')
            self.assertEqual(row['Künye (IEEE)'], 'M. Aksoy, “Benchmark RAG,” 2026.')
            self.assertEqual(row['Yıl · Tür · Kalite'], '2026 · Konferans · CORE A* (2023)')
            self.assertEqual(row['Okuyan üye · Düzey'], 'Hasan Basri Taşkın · 3. geçiş')
            self.assertFalse(row['DOI'])  # a Scopus record page is not a publisher link
            self.assertEqual([[c.value for c in r] for r in wb['Hariç'].iter_rows()],
                             [['Başlık', 'DOI', 'Gerekçe kodu'], ['Closed paper', '10.1007/x', 'H3']])
            text = ' '.join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value)
            self.assertNotIn('14.10.2026', text)
            self.assertNotIn(' ', text)
        finally:
            wb.close()
        again = self.store.export_delivery_workbook()
        self.assertEqual(again['relative_path'], result['relative_path'])
        self.assertEqual(sorted(p.name for p in Path(self.tmp.name, 'p', '05_Teslim').iterdir()), ['OEC-1017_O2_LiteraturMatrisi.xlsx'])


if __name__ == '__main__':
    unittest.main()
